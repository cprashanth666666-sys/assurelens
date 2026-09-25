"""Extracts plain text from a document's raw bytes, by sniffed MIME type.

Each extractor is defensive: a corrupt or hostile file must produce a failed
extraction, recorded as such, never an unhandled exception that takes the
background job down with it. A document that cannot be read is evidence of
that fact (`ExtractionResult.error`), not a crash.
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass

from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from openpyxl import load_workbook
from pypdf import PdfReader

from app.ingest.sniff import sniff_mime


@dataclass
class ExtractionResult:
    text: str | None
    page_count: int | None
    engine: str
    error: str | None = None


def extract_text(data: bytes, detected_mime: str | None = None) -> ExtractionResult:
    mime = detected_mime or sniff_mime(data)

    try:
        if mime == "application/pdf":
            return _extract_pdf(data)
        if mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            return _extract_docx(data)
        if mime == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
            return _extract_xlsx(data)
        if mime == "text/html":
            return _extract_html(data)
        if mime == "text/plain":
            return _extract_plain_or_csv(data)
    except Exception as exc:  # noqa: BLE001 — a bad file must not crash the pipeline
        return ExtractionResult(text=None, page_count=None, engine="none", error=str(exc))

    return ExtractionResult(
        text=None, page_count=None, engine="none",
        error=f"unsupported or undetected format (mime={mime!r})",
    )


def _extract_pdf(data: bytes) -> ExtractionResult:
    reader = PdfReader(io.BytesIO(data))
    pages = [page.extract_text() or "" for page in reader.pages]
    return ExtractionResult(
        text="\n\n".join(pages), page_count=len(reader.pages), engine="pypdf",
    )


def _extract_docx(data: bytes) -> ExtractionResult:
    doc = DocxDocument(io.BytesIO(data))
    paragraphs = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            paragraphs.append(" | ".join(cell.text for cell in row.cells))
    return ExtractionResult(text="\n".join(paragraphs), page_count=None, engine="python-docx")


def _extract_xlsx(data: bytes) -> ExtractionResult:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines: list[str] = []
    for sheet in wb.worksheets:
        lines.append(f"# {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                lines.append(" | ".join(cells))
    return ExtractionResult(
        text="\n".join(lines), page_count=len(wb.worksheets), engine="openpyxl",
    )


def _extract_html(data: bytes) -> ExtractionResult:
    soup = BeautifulSoup(data, "html.parser")
    for tag in soup(["script", "style"]):
        tag.decompose()
    text = soup.get_text(separator="\n")
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    return ExtractionResult(text="\n".join(lines), page_count=None, engine="beautifulsoup4")


def _extract_plain_or_csv(data: bytes) -> ExtractionResult:
    decoded = data.decode("utf-8", errors="replace")
    # A CSV renders as pipe-joined rows, same shape as the XLSX extractor,
    # so the classifier sees a consistent tabular text shape either way.
    sample = decoded[:2048]
    try:
        csv.Sniffer().sniff(sample)
        rows = list(csv.reader(io.StringIO(decoded)))
        text = "\n".join(" | ".join(row) for row in rows)
        return ExtractionResult(text=text, page_count=None, engine="csv")
    except csv.Error:
        return ExtractionResult(text=decoded, page_count=None, engine="plain-text")
