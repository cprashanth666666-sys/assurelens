"""The one-page executive summary. [PRD 7.2, S5]

Built from the same computations the in-product overview and roadmap use --
`get_readiness` and `get_roadmap` are called directly, not re-derived, so
the number on this page and the number on screen can never drift apart.
No headline compliance percentage, the same product-wide rule as everywhere
else. [PRD 7.3]
"""

from __future__ import annotations

import io

from docx import Document
from docx.document import Document as DocumentT
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app.api.controls import _latest_results, get_readiness
from app.api.findings import get_roadmap
from app.models.control import Control, Engagement
from app.reporting._docx import apply_base_styles, to_utc

DOMAIN_LABEL = {
    "NOTICE_CONSENT": "Notice & consent",
    "SECURITY_SAFEGUARDS": "Security safeguards",
    "RETENTION_ERASURE": "Retention & erasure",
    "PRINCIPAL_RIGHTS": "Principal rights",
    "THIRD_PARTY_TRANSFER": "Third party & transfer",
    "AI_GOVERNANCE": "AI governance",
    "BREACH_RESPONSE": "Breach response",
    "GOVERNANCE_ACCOUNTABILITY": "Governance & accountability",
}


def build_summary(db: Session, engagement_id: int) -> bytes:
    engagement = db.get(
        Engagement, engagement_id, options=[selectinload(Engagement.organization)]
    )
    if engagement is None:
        raise ValueError(f"No engagement with id {engagement_id}")

    readiness = get_readiness(engagement_id, db)
    roadmap = get_roadmap(engagement_id, db)
    latest = _latest_results(db, engagement_id)
    controls_by_id = {
        c.id: c for c in db.scalars(select(Control)).all()
    }
    gated = sorted(
        (
            (controls_by_id[cid], r)
            for cid, r in latest.items()
            if r.verdict == "INSUFFICIENT_EVIDENCE" and cid in controls_by_id
        ),
        key=lambda pair: pair[0].ref,
    )

    doc = Document()
    _set_core_properties(doc, engagement)
    apply_base_styles(doc, {1: 20, 2: 13.5})

    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph(engagement.name)
    p = doc.add_paragraph()
    p.add_run(engagement.organization.name).bold = True
    if engagement.organization.city:
        p.add_run(f" · {engagement.organization.city}")

    doc.add_heading("Where we stand, by domain", level=2)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    columns = ("Domain", "Pass", "Fail", "Insufficient evidence")
    for cell, text_ in zip(table.rows[0].cells, columns, strict=True):
        cell.text = text_
        for para in cell.paragraphs:
            for run in para.runs:
                run.bold = True
    for row in readiness.by_domain:
        cells = table.add_row().cells
        cells[0].text = DOMAIN_LABEL.get(row.domain, row.domain)
        cells[1].text = str(row.pass_)
        cells[2].text = str(row.fail)
        cells[3].text = str(row.insufficient_evidence)
    doc.add_paragraph().add_run(
        "Coverage is reported per domain, not as a single number: a domain "
        "at 40% coverage and one at 95% cannot be averaged into an honest "
        "figure."
    ).italic = True

    doc.add_heading("The five things to fix first", level=2)
    if roadmap.items:
        for item in roadmap.items[:5]:
            entry = doc.add_paragraph(style="List Number")
            entry.add_run(item.action)
            entry.add_run(f" (owner: {item.owner or 'unassigned'})").italic = True
    else:
        doc.add_paragraph("No open findings to sequence.")

    doc.add_heading("What we could not conclude, and why", level=2)
    if gated:
        for control, result in gated:
            entry = doc.add_paragraph(style="List Bullet")
            entry.add_run(f"{control.ref} — {control.title}: ").bold = True
            entry.add_run(
                result.gate_explanations[0] if result.gate_explanations
                else "Evidence did not clear the sufficiency gate."
            )
    else:
        doc.add_paragraph("Every tested control reached a Pass or Fail verdict.")

    doc.add_heading("What changes before the deadline", level=2)
    deadline = (
        f"{engagement.compliance_deadline:%d %B %Y}"
        if engagement.compliance_deadline else "the notified commencement date"
    )
    open_count = len(roadmap.items)
    doc.add_paragraph(
        f"Obligations commence {deadline}. {open_count} open finding"
        f"{'s' if open_count != 1 else ''} remain sequenced on the roadmap; "
        f"{len(gated)} control{'s' if len(gated) != 1 else ''} could not be "
        "concluded and need evidence before a verdict can be reported."
    )

    note = doc.add_paragraph()
    note.add_run(
        "Synthetic demonstration data. Not legal advice. Clause mappings "
        "are the author's reading of the DPDP Rules 2025."
    ).italic = True

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _set_core_properties(doc: DocumentT, engagement: Engagement) -> None:
    props = doc.core_properties
    props.title = f"AssureLens Executive Summary — {engagement.name}"
    props.author = "AssureLens"
    props.created = to_utc(engagement.created_at)
    props.modified = to_utc(engagement.created_at)
    props.revision = 1
