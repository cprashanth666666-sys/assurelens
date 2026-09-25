"""The DOCX workpaper. [Day 9, PRD 7.1]

Built against a real run on the seeded estate: python-docx round-trips the
result, so a structural assertion here (a heading exists, a table has the
right shape) also proves the file is well-formed OOXML that Word would open.
"""

from __future__ import annotations

import io

import pytest
from docx import Document
from sqlalchemy import select
from sqlalchemy.orm import Session

from app import suites as _suites  # noqa: F401  registers procedures
from app.engine.collectors import default_broker
from app.engine.runner import run_suites
from app.models.control import Engagement
from app.reporting.workpaper import LIMITATIONS, NOT_LEGAL_ADVICE, build_workpaper

pytestmark = pytest.mark.usefixtures("seeded_estate")


def _run(seeded_estate: Session, engagement_id: int) -> int:
    run = run_suites(
        seeded_estate, engagement_id, ["pii_retention", "third_party"],
        seed=42, broker=default_broker(), engine_version="test",
    )
    seeded_estate.flush()
    return run.id


def test_an_unknown_run_id_raises(seeded_estate: Session) -> None:
    with pytest.raises(ValueError, match="No run with id"):
        build_workpaper(seeded_estate, 999_999)
    seeded_estate.rollback()


def test_the_workpaper_opens_and_has_the_prescribed_sections(
    seeded_estate: Session,
) -> None:
    engagement = seeded_estate.scalar(select(Engagement).order_by(Engagement.id).limit(1))
    assert engagement is not None
    run_id = _run(seeded_estate, engagement.id)

    content = build_workpaper(seeded_estate, run_id)
    doc = Document(io.BytesIO(content))
    headings = [
        p.text for p in doc.paragraphs
        if p.style is not None and p.style.name.startswith("Heading")
    ]

    assert "Assurance Workpaper" in headings
    assert "Scope" in headings
    assert "Basis of preparation" in headings
    assert "Methodology" in headings
    assert "Limitations" in headings
    assert "Appendix A — Control library" in headings
    assert "Appendix B — Threshold register" in headings

    body = "\n".join(p.text for p in doc.paragraphs)
    assert NOT_LEGAL_ADVICE in body
    for item in LIMITATIONS:
        assert item in body

    # One control library row per authored control, plus the header row.
    library = doc.tables[0]
    assert len(library.rows) > 1
    assert [c.text for c in library.rows[0].cells] == [
        "Ref", "Title", "Domain", "Legal basis", "Executable",
    ]

    seeded_estate.rollback()


def test_a_control_section_names_its_result_and_prints_the_gate_reason(
    seeded_estate: Session,
) -> None:
    """DPDP-08-02 (erasure cascade) gates on G4 in the seeded estate -- the
    workpaper's "Why / Resolve" text must survive into the exported file,
    not just the on-screen render."""
    engagement = seeded_estate.scalar(select(Engagement).order_by(Engagement.id).limit(1))
    assert engagement is not None
    run_id = _run(seeded_estate, engagement.id)

    content = build_workpaper(seeded_estate, run_id)
    doc = Document(io.BytesIO(content))
    body = "\n".join(p.text for p in doc.paragraphs)

    assert "DPDP-08-02" in body
    assert "Insufficient Evidence" in body
    assert "Why." in body
    assert "Resolve." in body

    seeded_estate.rollback()


def test_a_not_applicable_control_does_not_describe_a_procedure_that_never_ran(
    seeded_estate: Session,
) -> None:
    """DPDP-08-04 does not bind this fictional BFSI GCC and is reported
    NOT_APPLICABLE. Its workpaper section must say so and stop -- not
    print "Procedure performed" / "Population and sample" text as if the
    control had actually been tested."""
    engagement = seeded_estate.scalar(select(Engagement).order_by(Engagement.id).limit(1))
    assert engagement is not None
    run_id = _run(seeded_estate, engagement.id)

    content = build_workpaper(seeded_estate, run_id)
    doc = Document(io.BytesIO(content))

    section_start = next(
        i for i, p in enumerate(doc.paragraphs) if p.text.startswith("DPDP-08-04")
    )
    next_heading = next(
        i for i, p in enumerate(doc.paragraphs[section_start + 1:], start=section_start + 1)
        if p.style is not None and p.style.name == "Heading 2"
    )
    section_text = "\n".join(p.text for p in doc.paragraphs[section_start:next_heading])

    assert "Out of scope" in section_text
    assert "Third Schedule" in section_text
    assert "Procedure performed" not in section_text
    assert "Population and sample" not in section_text
    assert "Evidence obtained" not in section_text

    seeded_estate.rollback()


def test_the_same_run_renders_byte_identical_output(seeded_estate: Session) -> None:
    engagement = seeded_estate.scalar(select(Engagement).order_by(Engagement.id).limit(1))
    assert engagement is not None
    run_id = _run(seeded_estate, engagement.id)

    first = build_workpaper(seeded_estate, run_id)
    second = build_workpaper(seeded_estate, run_id)
    assert first == second

    seeded_estate.rollback()
