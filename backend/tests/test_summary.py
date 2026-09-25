"""The one-page executive summary. [Day 9, PRD 7.2]"""

from __future__ import annotations

import io

import pytest
from docx import Document
from sqlalchemy import select
from sqlalchemy.orm import Session

from app import suites as _suites  # noqa: F401  registers procedures
from app.engine.collectors import default_broker
from app.engine.runner import run_suites
from app.models.control import Engagement
from app.reporting.summary import build_summary

pytestmark = pytest.mark.usefixtures("seeded_estate")


def test_an_unknown_engagement_raises(seeded_estate: Session) -> None:
    with pytest.raises(ValueError, match="No engagement"):
        build_summary(seeded_estate, 999_999)
    seeded_estate.rollback()


def test_the_summary_has_the_prescribed_sections_and_reuses_live_data(
    seeded_estate: Session,
) -> None:
    engagement = seeded_estate.scalar(select(Engagement).order_by(Engagement.id).limit(1))
    assert engagement is not None
    run_suites(
        seeded_estate, engagement.id, ["third_party"],
        seed=42, broker=default_broker(), engine_version="test",
    )
    seeded_estate.flush()

    content = build_summary(seeded_estate, engagement.id)
    doc = Document(io.BytesIO(content))
    headings = [
        p.text for p in doc.paragraphs
        if p.style is not None and p.style.name.startswith("Heading")
    ]

    assert "Executive Summary" in headings
    assert "Where we stand, by domain" in headings
    assert "The five things to fix first" in headings
    assert "What we could not conclude, and why" in headings
    assert "What changes before the deadline" in headings

    body = "\n".join(p.text for p in doc.paragraphs)

    # The erasure cascade that gates on G4 in test_workpaper.py must show up
    # here too, reusing get_readiness/get_roadmap rather than a second
    # computation that could drift from the on-screen one.
    assert "DPDP-08-02" in body

    library = doc.tables[0]
    assert [c.text for c in library.rows[0].cells] == [
        "Domain", "Pass", "Fail", "Insufficient evidence",
    ]

    seeded_estate.rollback()
